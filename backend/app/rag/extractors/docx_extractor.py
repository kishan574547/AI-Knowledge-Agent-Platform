import io
import logging
from typing import List
from docx import Document as DocxDocument
from app.rag.extractors.base import BaseExtractor, ExtractedDocument, PageContent
from app.core.errors import FileValidationException

logger = logging.getLogger("rag.extractors.docx")


class DocxExtractor(BaseExtractor):
    def extract(self, file_bytes: bytes, filename: str) -> ExtractedDocument:
        if not file_bytes:
            raise FileValidationException(f"DOCX file '{filename}' is empty")

        try:
            doc = DocxDocument(io.BytesIO(file_bytes))
        except Exception as e:
            logger.warning("Failed to parse DOCX %s: %s", filename, type(e).__name__)
            raise FileValidationException(f"Corrupted or invalid DOCX format in '{filename}'")

        paragraphs_text: List[str] = []

        try:
            for p in doc.paragraphs:
                if p.text.strip():
                    paragraphs_text.append(p.text.strip())

            # Also extract table text
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs_text.append(row_text)
        except Exception as e:
            logger.warning("Error reading DOCX paragraphs in %s: %s", filename, type(e).__name__)
            raise FileValidationException(f"Failed to extract text from DOCX '{filename}'")

        combined_text = "\n\n".join(paragraphs_text)
        return ExtractedDocument(
            text=combined_text,
            pages=[PageContent(page=None, text=combined_text)],
            metadata={"paragraph_count": len(paragraphs_text), "extractor": "docx"},
        )
